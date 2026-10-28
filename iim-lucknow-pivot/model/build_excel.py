"""Excel twin of dream11_pivot_model.py: inputs, outputs and a Checks sheet that ties to the JSON."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
O = json.load(open("model_outputs.json", encoding="utf-8"))
wb = Workbook(); HDR = PatternFill("solid", fgColor="5E0A0E"); INP = PatternFill("solid", fgColor="FFF4CC"); OUT = PatternFill("solid", fgColor="D8F0E6")
WHITE = Font(bold=True, color="FFFFFF"); B = Font(bold=True); thin = Side(style="thin", color="999999")
def sheet(title, header, rows, fill=None):
    ws = wb.create_sheet(title); ws.append(header)
    for c in ws[1]: c.fill = HDR; c.font = WHITE
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if fill: c.fill = fill
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or "")) for c in col) + 2))
    ws.sheet_view.showGridLines = False; return ws
wb.remove(wb.active)
readme = wb.create_sheet("README")
for line in ["Dream11 PIVOT '26 model, Team IdeaForge", "Values are pasted from model_outputs.json (the Python model is the source of truth).",
             "Yellow = inputs (case facts, researched facts, assumptions). Green = outputs. Checks sheet must show TRUE everywhere.",
             f"Model run: {O['meta']['run']}  hash: {O['meta'].get('hash','')}"]:
    readme.append([line])
readme.column_dimensions["A"].width = 110
flat = []
for grp in ["F", "R", "A"]:
    for k, v in O["meta"][grp].items(): flat.append([grp, k, json.dumps(v) if isinstance(v, (dict, list)) else v])
sheet("Inputs", ["Group", "Key", "Value"], flat, INP)
sheet("Base_Segments", ["ID", "Segment", "Pre-ban (M)", "Stay %", "Left (M)", "Now (M)", "Share %"],
      [[r["id"], r["name"], r["preban_m"], r["stay_pct"], r["left_m"], r["now_m"], r["now_share_pct"]] for r in O["segments"]["rows"]], OUT)
sheet("Options", ["Option"] + O["options"]["criteria"] + ["Weighted"], [[k + ". " + v["name"]] + v["scores"] + [v["weighted"]] for k, v in O["options"]["options"].items()], OUT)
sheet("RICE", ["Feature", "Reach", "Impact", "Confidence", "Effort", "RICE", "Cut"], [[r["feature"], r["reach"], r["impact"], r["confidence"], r["effort"], r["score"], r["cut"]] for r in O["rice"]], OUT)
for sc in ["bear", "base", "bull"]:
    ws = sheet(f"Revenue_{sc}", ["Month", "NSM (M)", "MAU (M)", "Ads", "Sponsorship", "Pro", "Gifting", "FanCode", "Total", "Total (formula)"],
               [[r["month"], r["nsm_m"], r["mau_m"], r["ads"], r["sponsorship"], r["pro"], r["gifting"], r["fancode"], r["total"]] for r in O["revenue"][sc]["rows"]], OUT)
    for i in range(2, 14): ws.cell(row=i, column=10, value=f"=SUM(D{i}:H{i})")
    ws.cell(row=15, column=1, value="12 months").font = B
    for col in "DEFGHI": ws[f"{col}15"] = f"=SUM({col}2:{col}13)"
sheet("Stress", ["Deposit case", "Deposit ratio", "Net pre-2023", "GST 28%", "Net 28%", "GST 40%", "Net 40%", "Drop 28 %", "Drop 40 %", "Prize pool 28 %"],
      [[sc, O["stress"][sc]["deposit_ratio"], O["stress"][sc]["net_pre2023"], O["stress"][sc]["gst28"], O["stress"][sc]["net28"], O["stress"][sc]["gst40"],
        O["stress"][sc]["net40"], O["stress"][sc]["drop28_pct"], O["stress"][sc]["drop40_pct"], O["stress"][sc]["prize_pool_to_hold_margin_28"]] for sc in ["bear", "base", "bull"]], OUT)
se = O["sensitivity"]; sheet("Sensitivity", ["NSM baseline (M)"] + [f"eCPM {e}" for e in se["ecpm_rs"]], [[w] + row for w, row in zip(se["wtb0"], se["grid"])], OUT)
ck = wb.create_sheet("Checks"); ck.append(["Check", "Result"])
for c in ck[1]: c.fill = HDR; c.font = WHITE
checks = [("Base 12-month total ties to JSON", f"=ABS(Revenue_base!I15-{O['revenue']['base']['total']['total']})<=0.5"),
          ("Bear 12-month total ties", f"=ABS(Revenue_bear!I15-{O['revenue']['bear']['total']['total']})<=0.5"),
          ("Bull 12-month total ties", f"=ABS(Revenue_bull!I15-{O['revenue']['bull']['total']['total']})<=0.5"),
          ("Monthly totals equal stream sums (base)", "=SUMPRODUCT(ABS(Revenue_base!I2:I13-Revenue_base!J2:J13))<=0.2"),
          ("Option D ranks first", f"=Options!H5=MAX(Options!H2:H6)"),
          ("Stress mid net at 28% below pre-2023", "=Stress!E3<Stress!C3"),
          ("Segments now sum to mid-2026 MAU", f"=ABS(SUM(Base_Segments!F2:F6)-{O['base']['mau_mid26_m']})<=0.3")]
for n, f in checks: ck.append([n, f])
ck.append(["All checks pass", f"=AND(B2:B{len(checks)+1})"])
ck.column_dimensions["A"].width = 50; ck.column_dimensions["B"].width = 16
wb.calculation.fullCalcOnLoad = True
wb.save("Dream11_Pivot_Model.xlsx"); print("saved")
