"""Build Eurogrip_Model.xlsx from model_outputs.json, model_tables.md and inputs/assumptions.csv.
  PY eurogrip-case/model/build_excel.py
Sheets: README, Assumptions, Case4_Summary, Case4_BreakEven_Live (live Excel formulas that rebuild the break-even
and check it against the Python model), T1 to T20 (every model table), Risks, Checks. No em or en dashes."""
import csv, io, json, pathlib, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

M = pathlib.Path(__file__).resolve().parent
O = json.load(io.open(M / "model_outputs.json", encoding="utf-8"))
SC = ("bear", "base", "bull")
HEADF = PatternFill("solid", fgColor="DCE4F0"); INF = PatternFill("solid", fgColor="FFF4CC")
HB = Font(bold=True, color="111111"); TF = Font(bold=True, size=16, color="184898"); NB = Font(bold=True)
thin = Side(style="thin", color="9CA3AF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); wb.remove(wb.active)


def sheet(name, title):
    ws = wb.create_sheet(name); ws["A1"] = title; ws["A1"].font = TF; return ws


def table(ws, r0, heads, rows):
    for j, h in enumerate(heads, 1):
        c = ws.cell(r0, j, h); c.font = HB; c.fill = HEADF; c.border = BOX
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.border = BOX
    return r0 + len(rows)


def widths(ws, ws_):
    for i, w in enumerate(ws_, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def num(s):
    t = str(s).strip().replace(",", "").replace("₹", "").replace("−", "-")
    return float(t) if re.fullmatch(r"-?\d+(\.\d+)?", t) else s


c4 = O["c4"]; A = O["A"]

# README
ws = sheet("README", "Eurogrip Case Study 4: model workbook")
lines = ["Team IdeaForge: Harsh Kawatra, Delhi Technological University.",
         "Prepared for the Eurogrip x Unstop Collaboration, Case Study 4.",
         f"Single source of truth: eurogrip_model.py. Model hash {O['meta']['hash']}, run {O['meta']['run_at']}.",
         "This workbook renders the model's outputs. One sheet, Case4_BreakEven_Live, rebuilds the break-even with live",
         "Excel formulas from editable inputs (yellow cells) and checks itself against the Python result.",
         "", "Sheet guide:",
         "Assumptions: all 130 assumptions, bear, base and bull, with units, basis and source IDs.",
         "Case4_Summary: the Case 4 headline outputs by scenario.",
         "Case4_BreakEven_Live: change any yellow cell to see the break-even move.",
         "T1 to T20: every model table (T1 to T12 Case 4, T13 to T15 Case 1, T16 and T17 Case 2, T18 to T20 Case 3).",
         "Risks: the Case 4 risk register with probability, impact and control.",
         "Checks: the model's own internal checks, all of which must read TRUE.",
         "", "Bear is the value least favourable to the recommendation; bull the most favourable."]
for i, line in enumerate(lines, 3):
    ws.cell(i, 1, line)
widths(ws, [120])

# Assumptions
ws = sheet("Assumptions", "Assumptions register (bear, base, bull)")
rows = []
for r in csv.DictReader(io.open(M / "inputs" / "assumptions.csv", encoding="utf-8-sig")):
    rows.append([r["id"], r["case"], r["label"], num(r["bear"]), num(r["base"]), num(r["bull"]), r["unit"], r["basis"], r["source_ids"]])
table(ws, 3, ["ID", "Case", "Assumption", "Bear", "Base", "Bull", "Unit", "Basis", "Sources"], rows)
widths(ws, [9, 6, 58, 11, 11, 11, 14, 90, 22])

# Case4_Summary
ws = sheet("Case4_Summary", "Case 4 headline outputs by scenario")
S = lambda k: [c4[k][s] for s in SC]
rows = [["Relevant unit revenue, ₹ crore"] + S("bu_rev_cr"),
        ["Budget share of unit revenue, %"] + [A["A02"][s] for s in SC],
        ["Budget envelope, ₹ crore"] + S("envelope_cr"),
        ["Programme spend, ₹ crore"] + S("spend_cr"),
        ["Spend as % of unit revenue"] + S("spend_share_pct"),
        ["Member shops, lakh"] + S("members_lakh"),
        ["Cost per member a year, ₹"] + S("cost_per_member"),
        ["Contribution per tyre, ₹"] + S("cpt"),
        ["Break-even extra tyres per member shop a year"] + S("be"),
        ["One extra tyre every (weeks)"] + S("weeks")]
rows = [[r[0]] + [round(v, 3) for v in r[1:]] for r in rows]
end = table(ws, 3, ["Output", "Bear", "Base", "Bull"], rows)
cs = c4["cash_scheme"]
table(ws, end + 3, ["The obvious cash-per-tyre scheme (base)", "Value"],
      [["Annual cost, ₹ crore", round(cs["cost_cr"], 2)], ["Share of spend on tyres that sell anyway, %", round(cs["deadweight_pct"], 1)],
       ["Tyres rewarded, lakh", round(cs["rewarded_units_lakh"], 2)]])
widths(ws, [50, 14, 14, 14])

# Case4_BreakEven_Live (live formulas)
ws = sheet("Case4_BreakEven_Live", "Break-even, rebuilt with live formulas (edit the yellow cells)")
heads = ["Item", "Bear", "Base", "Bull", "How it is computed"]
for j, h in enumerate(heads, 1):
    c = ws.cell(3, j, h); c.font = HB; c.fill = HEADF; c.border = BOX
inputs = [("Weighted replacement tyre MRP, ₹", "A27", "assumption A27"), ("GST on two-wheeler tyres, %", "A28", "assumption A28"),
          ("Channel margin on ex-GST price, %", "A29", "assumption A29"), ("Contribution margin, %", "A30", "assumption A30")]
r = 4
for label, k, how in inputs:
    ws.cell(r, 1, label).border = BOX
    for j, s in enumerate(SC, 2):
        c = ws.cell(r, j, A[k][s]); c.fill = INF; c.border = BOX
    ws.cell(r, 5, how).border = BOX; r += 1
for label, k, how in [("Programme spend, ₹ crore", "spend_cr", "model output c4.spend_cr"), ("Member shops, lakh", "members_lakh", "model output c4.members_lakh")]:
    ws.cell(r, 1, label).border = BOX
    for j, s in enumerate(SC, 2):
        c = ws.cell(r, j, round(c4[k][s], 6)); c.fill = INF; c.border = BOX
    ws.cell(r, 5, how).border = BOX; r += 1
# r is now 10
f = [("Net realisation per tyre, ₹", "={c}4/(1+{c}5/100)*(1-{c}6/100)", "MRP / (1 + GST) x (1 - channel margin)"),
     ("Contribution per tyre, ₹", "={c}10*{c}7/100", "net realisation x contribution margin"),
     ("Break-even extra tyres per shop a year", "=({c}8*10000000)/({c}9*100000)/{c}11", "spend / members / contribution per tyre"),
     ("One extra tyre every (weeks)", "=52/{c}12", "52 / break-even"),
     ("Break-even from the Python model", None, "model output c4.be"),
     ("Check", '=IF(ABS({c}12-{c}14)<0.01,"MATCHES MODEL","CHECK")', "live formula vs model")]
for i, (label, expr, how) in enumerate(f):
    rr = 10 + i
    ws.cell(rr, 1, label).font = NB; ws.cell(rr, 1).border = BOX
    for j, s in enumerate(SC, 2):
        col = get_column_letter(j)
        v = round(c4["be"][s], 6) if expr is None else expr.format(c=col)
        c = ws.cell(rr, j, v); c.border = BOX
    ws.cell(rr, 5, how).border = BOX
widths(ws, [44, 14, 14, 14, 44])
# Python replica of the formulas, so the workbook is proven before Excel ever opens it
for s in SC:
    net = A["A27"][s] / (1 + A["A28"][s] / 100) * (1 - A["A29"][s] / 100); cpt = net * A["A30"][s] / 100
    be = c4["spend_cr"][s] * 1e7 / (c4["members_lakh"][s] * 1e5) / cpt
    assert abs(be - c4["be"][s]) < 0.01, (s, be, c4["be"][s])

# T1 to T20 from model_tables.md
md = io.open(M / "model_tables.md", encoding="utf-8").read()
for block in re.split(r"\n(?=## T\d+\. )", md)[1:]:
    lines_ = block.strip().splitlines(); title = lines_[0][3:].strip(); tid = title.split(".")[0]
    tab = [l for l in lines_ if l.startswith("|")]
    src = next((l.strip("*") for l in lines_ if l.startswith("*Source")), "")
    cells = [[c.strip() for c in l.strip("|").split("|")] for l in tab]
    heads, body = cells[0], [[num(c) for c in row] for row in cells[2:]]
    ws = sheet(tid, title)
    end = table(ws, 3, heads, body)
    ws.cell(end + 2, 1, src)
    widths(ws, [48] + [16] * (len(heads) - 1))

# Risks
ws = sheet("Risks", "Case 4 risk register (probability and impact on a 1 to 3 scale)")
table(ws, 3, ["ID", "Risk", "P", "I", "P x I", "Control"],
      [[x["id"], x["risk_long"], x["p"], x["i"], x["p"] * x["i"], x["control_long"]] for x in c4["risks"]])
widths(ws, [6, 70, 5, 5, 7, 90])

# Checks
ws = sheet("Checks", "Model internal checks")
table(ws, 3, ["Check", "Result"], [[k, bool(v)] for k, v in O["checks"].items()] + [[f"flag: {k}", bool(v)] for k, v in O["flags"].items()])
widths(ws, [48, 12])

out = M / "Eurogrip_Model.xlsx"; wb.save(out); print("wrote", out)
