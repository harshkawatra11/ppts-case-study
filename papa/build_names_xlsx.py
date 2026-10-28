"""Build handwritten_names.xlsx from names transcribed out of the two
papa/ ledger-page images. Only handwritten names are included; printed
account/customer labels on the source pages are excluded."""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

names = [
    # --- Page 1 (WhatsApp Image ...20.21.572.jpeg) ---
    # Top-margin handwritten scratch list
    "Sunil Sharma",
    "Vijay Nagar",
    "Manish Singh",
    "Vinod Kalra",
    "Arpit Kalra",
    "Nikunj",
    "Nikky",
    "Naveen Goyal",
    "Vinit Gupta",
    "Doulat Jain",
    "Harish Goel",
    "Shivam Goel",
    # Just above the table header
    "Shresth Aggarwal",
    "Arun Aggarwal",
    # Table rows, top to bottom
    "Abhay Kumar",
    "Abhishek",
    "Ankit",
    "Lokesh Gupta",
    "Amirchand",
    "Ayush Gupta",
    "Divansh",
    "Amit Sharma",
    "Jatan Gupta",
    "Abhishek Gupta",
    "Anuj Jain",
    "Manoj Gupta",
    "Sachin Jain",
    "Ashu Juneja",
    "Raman Juneja",
    "Vishal Tayal",
    "Hariom Tayal",
    "Priyank Tayal",
    "R.S. Sharma",
    "Nittin Sharma",
    "Shubham",
    "Krishan Jindal",
    "Rasu Manchanda",
    "Sunil Kapoor",
    "Ashok Chawla",
    "KH",
    "Sunil Tayal",
    "Vinay Goyal",
    "Mahender",
    "Aniket Bansal",
    "Mukesh Grover",
    "Sanjay Aggarwal",
    "Manoj Aggarwal",
    "Sunil Aggarwal",
    "Manoj Bansal",
    "Kamal Aggarwal",
    "Satish Jain",
    "Vijay Juneja",
    "Vippin Jain",
    "Ajay Jain",
    "Manoj Sharma",
    "Pankaj",
    "Manoj",

    # --- Page 2 (WhatsApp Image ...20.21.57.jpeg) ---
    "Dharampal Gupta",
    "Mahaveer Daasa",
    "Navratan Daasa",
    "Neeraj Goyal",
    "Neeraj Tayal",
    "Arun Varma",
    "Om Parkash",
    "Pappu Ji",
    "Ravinder Sharma",
    "Dilip Jain",
    "Poonam Jain",
    "Rahul Sharma",
    "Anurag Goyal",
    "Rakesh Gupta",
    "Sanjay Gupta",
    "Moolchand Gupta",
    "Deepak Ji",
    "Shri Parkash",
    "Rommy Ji",
    "SK Garg",
    "Shantanu Gupta",
    "Shivom Singhal",
    "Sukhraj",
    "Rajeev Surana",
    "Shivkumar Gupta",
    "Tinku",
    "Kashmiri",
]

wb = Workbook()
ws = wb.active
ws.title = "Handwritten Names"

ws["A1"] = "S.No."
ws["B1"] = "Name"
for cell in ("A1", "B1"):
    ws[cell].font = Font(bold=True, size=12)
    ws[cell].alignment = Alignment(horizontal="center")

for i, name in enumerate(names, start=1):
    ws.cell(row=i + 1, column=1, value=i).alignment = Alignment(horizontal="center")
    ws.cell(row=i + 1, column=2, value=name)

ws.column_dimensions[get_column_letter(1)].width = 10
ws.column_dimensions[get_column_letter(2)].width = 30
ws.freeze_panes = "A2"

out_path = "papa/handwritten_names.xlsx"
wb.save(out_path)
print(f"Saved {len(names)} names to {out_path}")
