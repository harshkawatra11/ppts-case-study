# ===== SETU WORKBOOK =====
"""Writes model/Setu_Model.xlsx: every table of model_tables.md as a sheet, plus an Inputs sheet (case facts, assumptions, levers).
Values workbook, not a formula twin. Run after make_values.py:  python lsr-caseforchange/model/build_excel.py"""
import json, pathlib, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
HERE = pathlib.Path(__file__).resolve().parent; O = json.loads((HERE / "model_outputs.json").read_text(encoding="utf-8"))
wb = Workbook(); ws = wb.active; ws.title = "Inputs"; HEAD = PatternFill("solid", fgColor="0B1437"); WHITE = Font(bold=True, color="FFFFFF")
def put(ws, r, vals, head=False):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        if head: cell.fill = HEAD; cell.font = WHITE; cell.alignment = Alignment(wrap_text=True, vertical="top")
def flat(d, pre=""):
    for k, v in d.items():
        if isinstance(v, dict) and not all(str(x).isdigit() for x in v): yield from flat(v, pre + k + ".")
        else: yield pre + k, json.dumps(v) if isinstance(v, (dict, list)) else v
put(ws, 1, ["Block", "Key", "Value"], True); r = 2
for blk, name in (("F", "Case fact"), ("A", "Team assumption"), ("L", "Plan lever"), ("SCEN", "Scenario override"), ("CONT", "Contingency override")):
    for k, v in flat(O["meta"][blk]): put(ws, r, [name, k, v]); r += 1
put(ws, r + 1, ["Model hash", O["meta"]["hash"]])
blocks = re.split(r"\n## ", (HERE / "model_tables.md").read_text(encoding="utf-8"))[1:]
for b in blocks:
    title = b.split("\n", 1)[0]; ws = wb.create_sheet(re.sub(r"[^A-Za-z0-9 ]", "", title)[:28]); put(ws, 1, [title]); rr = 3; first = True
    for line in b.split("\n")[1:]:
        if not line.startswith("|"): first = True; rr += 0 if first else 1; continue
        if set(line.replace("|", "").strip()) <= set("-"): continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        vals = []
        for c in cells:
            x = c.replace("−", "-").replace(",", "")
            try: vals.append(float(x) if "." in x else int(x))
            except ValueError: vals.append(c)
        put(ws, rr, vals if not first else cells, first); first = False; rr += 1
for ws in wb.worksheets:
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 22
wb.save(HERE / "Setu_Model.xlsx"); print("sheets", len(wb.worksheets))
